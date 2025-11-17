# main.py
import os
import csv
from datetime import datetime
from functools import partial

from openpyxl import Workbook  # Excel export (Android compatible)

from kivy.clock import Clock
from kivy.lang import Builder
from kivy.metrics import dp
from kivy.properties import ListProperty, StringProperty
from kivymd.app import MDApp
from kivymd.uix.boxlayout import MDBoxLayout
from kivymd.uix.dialog import MDDialog
from kivymd.uix.menu import MDDropdownMenu
from kivymd.uix.list import OneLineListItem
from kivymd.uix.pickers import MDDatePicker
from kivymd.uix.datatables import MDDataTable

# -------------------------
# Paths
# -------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MAPPINGS_CSV = os.path.join(BASE_DIR, "mappings.csv")
FAULT_TYPES_CSV = os.path.join(BASE_DIR, "fault_types.csv")
DAILY_LOG_CSV = os.path.join(BASE_DIR, "daily_log.csv")

# -------------------------
# Utility: Load CSV → list of dicts
# -------------------------
def load_csv_to_dicts(path):
    rows = []
    if not os.path.exists(path):
        return rows
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append(r)
    return rows


# Load CSVs
m_df = load_csv_to_dicts(MAPPINGS_CSV)
ft_df = load_csv_to_dicts(FAULT_TYPES_CSV)

# -------------------------
# FAULT TYPES LIST
# -------------------------
FAULT_TYPES = [row["Type of Fault"] for row in ft_df if row.get("Type of Fault")]

# -------------------------
# Default Excel export path
# -------------------------
if "ANDROID_ARGUMENT" in os.environ:
    DEFAULT_EXPORT = "/sdcard/Download/logbook_data.xlsx"
else:
    DEFAULT_EXPORT = os.path.join(BASE_DIR, "logbook_data.xlsx")

# -------------------------
# Build cascading mapping structure
# -------------------------
mapping = {}

for r in m_df:
    line = str(r.get("Line", "")).strip()
    area = str(r.get("Area", "")).strip()
    etype = str(r.get("Equipment Type", "")).strip()
    equip = str(r.get("Equipment", "")).strip()

    if not line:
        continue

    mapping.setdefault(line, {}).setdefault(area, {}).setdefault(etype, set()).add(equip)

# convert sets → sorted lists
for line in mapping:
    for area in mapping[line]:
        for etype in mapping[line][area]:
            mapping[line][area][etype] = sorted(list(mapping[line][area][etype]))

# -------------------------
# Ensure daily_log.csv exists
# -------------------------
if not os.path.exists(DAILY_LOG_CSV):
    with open(DAILY_LOG_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Time", "Date", "Line", "Area", "Equipment_Type", "Equipment",
            "Problem_Description", "Action_Type", "Action", "Type_of_Fault",
            "Job_Type", "EST", "LOTO_Applied"
        ])


# =======================================================================
#                           MAIN FORM CLASS
# =======================================================================

class LogForm(MDBoxLayout):
    lines = ListProperty(sorted(list(mapping.keys())))
    areas = ListProperty([])
    types = ListProperty([])
    equips = ListProperty([])
    fault_types = ListProperty(FAULT_TYPES)
    status = StringProperty("Ready")

    def __init__(self, **kwargs):
        super().__init__(**kwargs)

        self._preview_dialog = None

        # Load saved entries
        try:
            self.saved_entries = load_csv_to_dicts(DAILY_LOG_CSV)
        except:
            self.saved_entries = []

        Clock.schedule_once(self._fill_datetime, 0.1)

        # dropdown menu holders
        self._menus = {}

        # static dropdown options
        self._static_options = {
            "job_type": ["Breakdown Maint", "Preventive Maint", "Operational Maint"],
            "action_type": ["Reset/Adjustment", "Repair/Replacement"],
            "loto": ["Yes", "No"],
        }

        Clock.schedule_once(lambda dt: self._prepare_static_menus(), 0.2)

    def _fill_datetime(self, dt=None):
        now = datetime.now()
        self.ids.date_input.text = now.strftime("%d/%m/%Y")
        self.ids.time_input.text = now.strftime("%H:%M:%S")

    # ----------------------------------------------------------
    # Date Picker
    # ----------------------------------------------------------
    def open_date_picker(self):
        date_dialog = MDDatePicker(year=datetime.now().year,
                                   month=datetime.now().month,
                                   day=datetime.now().day)
        date_dialog.bind(on_save=self._on_date_save)
        date_dialog.open()

    def _on_date_save(self, instance, value, range_range):
        self.ids.date_input.text = value.isoformat()

    # ----------------------------------------------------------
    # Dropdown Menus (searchable dynamic)
    # ----------------------------------------------------------
    def _prepare_static_menus(self):
        for key, options in self._static_options.items():
            field = f"{key}_field"
            self._create_menu(field, options, searchable=False)

        # create menus for cascading dropdowns
        self._create_menu("line_field", self.lines, searchable=True)
        self._create_menu("area_field", self.areas, searchable=True)
        self._create_menu("type_field", self.types, searchable=True)
        self._create_menu("equip_field", self.equips, searchable=True)
        self._create_menu("fault_type_field", self.fault_types, searchable=True)

    def _create_menu(self, field_id, options, searchable=False):
        old = self._menus.get(field_id)
        if old:
            try:
                old.dismiss()
            except:
                pass

        caller = self.ids.get(field_id)
        if not caller:
            return

        def make_items(opt_list):
            return [{
                "viewclass": "OneLineListItem",
                "text": str(opt),
                "height": dp(44),
                "on_release": partial(self._on_menu_select, field_id, str(opt))
            } for opt in opt_list]

        items = make_items(options)

        menu = MDDropdownMenu(
            caller=caller,
            items=items,
            width_mult=4,
            max_height=dp(300),
        )
        self._menus[field_id] = menu

        menu.is_open_flag = False
        menu.bind(on_open=lambda *x: setattr(menu, "is_open_flag", True))
        menu.bind(on_dismiss=lambda *x: setattr(menu, "is_open_flag", False))

        if searchable:
            caller.readonly = False

            def on_text(instance, value):
                filtered = [o for o in options if value.lower() in str(o).lower()]
                menu.items = make_items(filtered)
                if filtered and not menu.is_open_flag:
                    try:
                        menu.open()
                    except:
                        pass

            caller.bind(text=on_text)

    # handle dropdown select
    def _on_menu_select(self, field_id, selected_text, *args):
        fld = self.ids.get(field_id)
        if fld:
            fld.text = selected_text

        menu = self._menus.get(field_id)
        if menu:
            menu.dismiss()

        # cascading updates
        if field_id == "line_field":
            self.on_line_changed(selected_text)
        elif field_id == "area_field":
            self.on_area_changed(selected_text)
        elif field_id == "type_field":
            self.on_type_changed(selected_text)

    # ----------------------------------------------------------
    # Cascading Dropdown Logic
    # ----------------------------------------------------------
    def on_line_changed(self, value):
        self.areas = sorted(list(mapping.get(value, {}).keys()))
        self.ids.area_field.text = "Select Area"
        self._create_menu("area_field", self.areas, searchable=True)

        self.types = []
        self.equips = []
        self.ids.type_field.text = "Select Type"
        self.ids.equip_field.text = "Select Equipment"
        self._create_menu("type_field", [], searchable=True)
        self._create_menu("equip_field", [], searchable=True)

    def on_area_changed(self, value):
        line = self.ids.line_field.text
        if line in mapping and value in mapping[line]:
            self.types = sorted(list(mapping[line][value].keys()))
        else:
            self.types = []

        self.ids.type_field.text = "Select Type"
        self._create_menu("type_field", self.types, searchable=True)

        self.equips = []
        self.ids.equip_field.text = "Select Equipment"
        self._create_menu("equip_field", [], searchable=True)

    def on_type_changed(self, value):
        line = self.ids.line_field.text
        area = self.ids.area_field.text
        self.equips = mapping.get(line, {}).get(area, {}).get(value, [])
        self.ids.equip_field.text = "Select Equipment"
        self._create_menu("equip_field", self.equips, searchable=True)

    # ----------------------------------------------------------
    # Form Validation
    # ----------------------------------------------------------
    def validate_form(self):
        required = ["fault_input", "action_input", "est_input"]

        for fid in required:
            if not self.ids[fid].text.strip():
                return f"Please fill the '{fid.replace('_',' ').title()}' field."

        dropdowns = [
            "line_field", "area_field", "type_field", "equip_field",
            "fault_type_field", "job_type_field", "loto_field", "action_type_field",
        ]

        for fid in dropdowns:
            text = self.ids[fid].text.strip()
            if not text or text.startswith("Select"):
                return f"Please select a valid value for '{fid.replace('_',' ').title()}'"

        return None

    # ----------------------------------------------------------
    # Save Entry
    # ----------------------------------------------------------
    def save_entry(self):
        rec = {
            "Time": self.ids.time_input.text,
            "Date": self.ids.date_input.text,
            "Line": self.ids.line_field.text,
            "Area": self.ids.area_field.text,
            "Equipment_Type": self.ids.type_field.text,
            "Equipment": self.ids.equip_field.text,
            "Problem_Description": self.ids.fault_input.text,
            "Action_Type": self.ids.action_type_field.text,
            "Action": self.ids.action_input.text,
            "Type_of_Fault": self.ids.fault_type_field.text,
            "Job_Type": self.ids.job_type_field.text,
            "EST": self.ids.est_input.text,
            "LOTO_Applied": self.ids.loto_field.text
        }

        error = self.validate_form()
        if error:
            MDDialog(text=error, size_hint=(0.7, 0.3)).open()
            return

        self.saved_entries.append(rec)

        with open(DAILY_LOG_CSV, "a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=rec.keys())
            writer.writerow(rec)

        self.status = f"Saved ({len(self.saved_entries)} rows)"
        self._fill_datetime()
        self.clear_form()

    # ----------------------------------------------------------
    # Clear form fields
    # ----------------------------------------------------------
    def clear_form(self):
        for fid in ["fault_input", "action_input", "est_input"]:
            self.ids[fid].text = ""

        placeholders = {
            "line_field": "Select Line",
            "area_field": "Select Area",
            "type_field": "Select Type",
            "equip_field": "Select Equipment",
            "fault_type_field": "Select Fault Type",
            "job_type_field": "Select Job Type",
            "loto_field": "Select LOTO",
            "action_type_field": "Select Action Type",
        }

        for fid, placeholder in placeholders.items():
            self.ids[fid].text = placeholder

        self._prepare_static_menus()

    # ----------------------------------------------------------
    # Preview Entries
    # ----------------------------------------------------------
    def preview(self):
        if not self.saved_entries:
            MDDialog(text="No entries to preview", size_hint=(0.7, 0.3)).open()
            return

        # Close old one
        if self._preview_dialog:
            try: self._preview_dialog.dismiss()
            except: pass

        columns = [(col.replace("_"," "), dp(30)) for col in self.saved_entries[0].keys()]
        rows = [tuple(str(v) for v in row.values()) for row in self.saved_entries]

        table = MDDataTable(
            size_hint=(1, None),
            height=dp(600),
            column_data=columns,
            row_data=rows,
            check=False,
            use_pagination=False,
        )

        self._preview_dialog = MDDialog(
            title=f"Preview ({len(rows)} rows)",
            type="custom",
            content_cls=table,
            size_hint=(0.95, 0.95),
        )
        self._preview_dialog.open()

    # ----------------------------------------------------------
    # Export to Excel (openpyxl)
    # ----------------------------------------------------------
    def export_xlsx(self, path=None):
        if not self.saved_entries:
            self.status = "No entries to export"
            return

        if path is None:
            path = DEFAULT_EXPORT

        wb = Workbook()
        ws = wb.active

        headers = list(self.saved_entries[0].keys())
        ws.append(headers)

        for row in self.saved_entries:
            ws.append([row[h] for h in headers])

        wb.save(path)
        self.status = f"Exported {len(self.saved_entries)} rows → {path}"

        # reset CSV
        with open(DAILY_LOG_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

        self.saved_entries = []

    # ----------------------------------------------------------
    # Clear saved entries
    # ----------------------------------------------------------
    def clear_saved(self):
        self.saved_entries = []
        with open(DAILY_LOG_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([
                "Time", "Date", "Line", "Area", "Equipment_Type", "Equipment",
                "Problem_Description", "Action_Type", "Action", "Type_of_Fault",
                "Job_Type", "EST", "LOTO_Applied"
            ])
        self.status = "Cleared"

    # ----------------------------------------------------------
    # Theme toggle
    # ----------------------------------------------------------
    def toggle_theme(self):
        app = MDApp.get_running_app()
        app.theme_cls.theme_style = (
            "Light" if app.theme_cls.theme_style == "Dark" else "Dark"
        )
        for k, menu in self._menus.items():
            if menu:
                try:
                    menu.dismiss()
                except:
                    pass
        Clock.schedule_once(lambda dt: self._prepare_static_menus(), 0.05)


# ============================================================
# Load KV and run App
# ============================================================
KV = Builder.load_file(os.path.join(BASE_DIR, "logform3.kv"))

class LogApp(MDApp):
    def build(self):
        self.theme_cls.theme_style = "Light"
        self.theme_cls.primary_palette = "Teal"
        self.theme_cls.accent_palette = "Amber"
        return LogForm()

if __name__ == "__main__":
    LogApp().run()
